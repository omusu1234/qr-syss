from flask import Flask, render_template, request, jsonify, redirect, url_for, flash, send_file
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from config import Config
from models import db, User, Lecturer, Unit, LectureSession, Attendance, Notification, Department, ActivityLog, LoginHistory, PhotoMatch
from database_sync import DatabaseSync
from datetime import datetime, timedelta, timezone
from sqlalchemy import text
from collections import defaultdict
import qrcode
import io
import base64
import secrets
import csv
from geopy.distance import geodesic
from functools import wraps
import os
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

app = Flask(__name__)
app.config.from_object(Config)

# Configure logging to stdout (for Render and other cloud platforms)
if not app.debug:
    # In production, log to stdout so Render can capture it
    stream_handler = logging.StreamHandler()
    stream_handler.setLevel(logging.INFO)
    stream_handler.setFormatter(logging.Formatter(
        '[%(asctime)s] %(levelname)s in %(module)s: %(message)s'
    ))
    app.logger.addHandler(stream_handler)
    app.logger.setLevel(logging.INFO)
    app.logger.info('Application logging configured')

# Timezone configuration (EAT - East Africa Time, UTC+3)
TIMEZONE_OFFSET_HOURS = int(os.environ.get('TIMEZONE_OFFSET_HOURS', 3))  # Default to UTC+3 (EAT)

@app.template_filter('localtime')
def localtime_filter(dt):
    """Convert UTC datetime to local time (EAT - UTC+3 by default)"""
    if dt is None:
        return None
    if isinstance(dt, str):
        dt = datetime.fromisoformat(dt.replace('Z', '+00:00'))
    if dt.tzinfo is None:
        # Assume UTC if no timezone info
        dt = dt.replace(tzinfo=timezone.utc)
    # Convert to local time (UTC+3 for EAT)
    local_tz = timezone(timedelta(hours=TIMEZONE_OFFSET_HOURS))
    local_dt = dt.astimezone(local_tz)
    return local_dt

db.init_app(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Please log in to access this page.'

# Database sync will be initialized on first use
db_sync = None

def get_db_sync():
    """Get or create database sync instance"""
    global db_sync
    if db_sync is None:
        db_sync = DatabaseSync()
    return db_sync

def get_base_url():
    """Get the base URL for QR code generation (production or local)"""
    # Check for environment variable first (for production)
    base_url = os.environ.get('BASE_URL')
    if base_url:
        return base_url.rstrip('/')
    
    # For production on Render, always use HTTPS
    if request.host:
        # Check if we're on Render or any production environment
        if 'render.com' in request.host or 'onrender.com' in request.host:
            # Always use HTTPS for Render
            return f"https://{request.host}"
        # Check if request is already HTTPS
        if request.is_secure or request.headers.get('X-Forwarded-Proto') == 'https':
            return f"https://{request.host}"
    
    # For local development, use the configured IP
    local_ip = os.environ.get('LOCAL_IP', '10.10.185.117')
    port = request.environ.get('SERVER_PORT', '5002')
    return f"http://{local_ip}:{port}"

def get_client_ip():
    """Get client IP address, handling proxy headers"""
    ip = request.remote_addr
    if request.headers.get('X-Forwarded-For'):
        ip = request.headers.get('X-Forwarded-For').split(',')[0].strip()
    return ip

def log_activity(action, entity_type=None, entity_id=None, description=None, user_id=None):
    """Log user activity to audit trail"""
    try:
        log = ActivityLog(
            user_id=user_id or (current_user.id if current_user.is_authenticated else None),
            action=action,
            entity_type=entity_type,
            entity_id=entity_id,
            description=description,
            ip_address=get_client_ip()
        )
        db.session.add(log)
        db.session.commit()
    except Exception as e:
        app.logger.error(f"Failed to log activity: {str(e)}")

def log_login(user, success=True, failure_reason=None):
    """Log login attempt - only logs if user exists (user_id is required)"""
    try:
        # Only log if user exists (user_id is required in database)
        if not user:
            return  # Can't log without a user
        
        login_log = LoginHistory(
            user_id=user.id,
            ip_address=get_client_ip(),
            user_agent=request.headers.get('User-Agent'),
            success=success,
            failure_reason=failure_reason
        )
        db.session.add(login_log)
        db.session.commit()
    except Exception as e:
        app.logger.error(f"Failed to log login: {str(e)}", exc_info=True)

@login_manager.user_loader
def load_user(user_id):
    """Load user by ID - required by Flask-Login"""
    try:
        if user_id is None:
            return None
        user = User.query.get(int(user_id))
        return user
    except (ValueError, TypeError) as e:
        app.logger.warning(f"Invalid user_id format: {user_id}, error: {str(e)}")
        return None
    except Exception as e:
        app.logger.error(f"Error loading user {user_id}: {str(e)}", exc_info=True)
        return None

def admin_required(f):
    """Decorator to require admin role"""
    @wraps(f)
    @login_required
    def decorated_function(*args, **kwargs):
        if not current_user.is_admin():
            flash('Access denied. Admin privileges required.', 'error')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated_function

def lecturer_required(f):
    """Decorator to require lecturer role"""
    @wraps(f)
    @login_required
    def decorated_function(*args, **kwargs):
        if not current_user.is_lecturer():
            flash('Access denied. Lecturer privileges required.', 'error')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated_function

# Error handlers
@app.errorhandler(500)
def internal_error(error):
    """Handle internal server errors"""
    db.session.rollback()
    app.logger.error(f"Internal server error: {str(error)}", exc_info=True)
    return render_template('error.html', 
                         message=f'An internal server error occurred. Please try again. Error: {str(error)}'), 500

@app.errorhandler(404)
def not_found_error(error):
    """Handle 404 errors"""
    return render_template('error.html', message='Page not found'), 404

# Note: Removed global Exception handler as it was interfering with normal operations
# Flask's debug mode will show detailed errors in development

# Test route to check if server is working
@app.route('/health')
def health_check():
    """Health check endpoint"""
    try:
        # Test database connection
        db.session.execute(text('SELECT 1'))
        return jsonify({'status': 'ok', 'database': 'connected'}), 200
    except Exception as e:
        return jsonify({'status': 'error', 'database': 'disconnected', 'error': str(e)}), 500

# Debug endpoint for testing coordinate calculations
@app.route('/debug/distance', methods=['GET'])
def debug_distance():
    """Debug endpoint to test distance calculations between two coordinates"""
    try:
        lat1 = float(request.args.get('lat1', 0))
        lon1 = float(request.args.get('lon1', 0))
        lat2 = float(request.args.get('lat2', 0))
        lon2 = float(request.args.get('lon2', 0))
        
        # Calculate all combinations
        normal = geodesic((lat1, lon1), (lat2, lon2)).meters
        swap1 = geodesic((lat1, lon1), (lon2, lat2)).meters
        swap2 = geodesic((lon1, lat1), (lat2, lon2)).meters
        swap_both = geodesic((lon1, lat1), (lon2, lat2)).meters
        
        return jsonify({
            'point1': {'lat': lat1, 'lon': lon1},
            'point2': {'lat': lat2, 'lon': lon2},
            'distances': {
                'normal': normal,
                'swap_student': swap1,
                'swap_lecture': swap2,
                'swap_both': swap_both
            },
            'best': min(normal, swap1, swap2, swap_both)
        }), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 400

# Routes
@app.route('/')
def index():
    try:
        if current_user.is_authenticated:
            return redirect(url_for('dashboard'))
    except Exception as e:
        app.logger.error(f"Error checking authentication: {str(e)}")
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    try:
        if request.method == 'POST':
            try:
                username = request.form.get('username')
                password = request.form.get('password')
                role = request.form.get('role')  # 'lecturer' or 'admin'
                
                app.logger.info(f"Login attempt: username={username}, role={role}")
                
                if not username or not password or not role:
                    flash('Please fill in all fields', 'error')
                    return render_template('login.html')
                
                user = User.query.filter_by(username=username, role=role).first()
                
                if user:
                    app.logger.info(f"User found: {user.username}, checking password...")
                    if user.check_password(password):
                        app.logger.info(f"Password correct, logging in user: {user.username}")
                        login_user(user)
                        # Log login and activity (non-blocking - don't fail login if logging fails)
                        try:
                            log_login(user, success=True)
                            log_activity('login', description=f'User {user.username} logged in')
                        except Exception as e:
                            app.logger.error(f"Failed to log login activity: {str(e)}")
                        
                        app.logger.info(f"Redirecting to dashboard for user: {user.username}")
                        return redirect(url_for('dashboard'))
                    else:
                        app.logger.warning(f"Invalid password for user: {username}")
                        # Log failed login attempt (non-blocking)
                        try:
                            log_login(user, success=False, failure_reason='Invalid password')
                            log_activity('login_failed', description=f'Failed login attempt for {username}')
                        except Exception as e:
                            app.logger.error(f"Failed to log failed login: {str(e)}")
                        flash('Invalid username, password, or role', 'error')
                else:
                    app.logger.warning(f"User not found: username={username}, role={role}")
                    flash('Invalid username, password, or role', 'error')
            except Exception as e:
                app.logger.error(f"Login POST error: {str(e)}", exc_info=True)
                flash(f'An error occurred during login: {str(e)}', 'error')
        
        # GET request - render login page
        return render_template('login.html')
    except Exception as e:
        app.logger.error(f"Login route error: {str(e)}", exc_info=True)
        # Return a simple error page if template rendering fails
        return f"""
        <html>
        <head><title>Error</title></head>
        <body>
            <h1>Internal Server Error</h1>
            <p>An error occurred: {str(e)}</p>
            <p>Please check the server logs for more details.</p>
        </body>
        </html>
        """, 500

@app.route('/logout')
@login_required
def logout():
    # Log logout
    if current_user.is_authenticated:
        log_activity('logout', description=f'User {current_user.username} logged out')
        # Update last login history entry
        last_login = LoginHistory.query.filter_by(
            user_id=current_user.id,
            logout_at=None
        ).order_by(LoginHistory.login_at.desc()).first()
        if last_login:
            last_login.logout_at = datetime.utcnow()
            db.session.commit()
    logout_user()
    flash('You have been logged out', 'info')
    return redirect(url_for('login'))

@app.route('/dashboard')
@login_required
def dashboard():
    if current_user.is_admin():
        return redirect(url_for('admin_dashboard'))
    elif current_user.is_lecturer():
        return redirect(url_for('lecturer_dashboard'))
    return redirect(url_for('login'))

@app.route('/admin/dashboard')
@admin_required
def admin_dashboard():
    lecturers = Lecturer.query.all()
    units = Unit.query.all()
    # Get recent notifications count
    recent_notifications = Notification.query.order_by(Notification.created_at.desc()).limit(5).all()
    return render_template('admin_dashboard.html', 
                         lecturers=lecturers, 
                         units=units,
                         recent_notifications=recent_notifications)

@app.route('/lecturer/dashboard')
@lecturer_required
def lecturer_dashboard():
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    units = Unit.query.filter_by(lecturer_id=lecturer.id).all()
    recent_sessions = LectureSession.query.join(Unit).filter(
        Unit.lecturer_id == lecturer.id
    ).order_by(LectureSession.created_at.desc()).limit(5).all()
    
    # Get recent notifications for lecturers
    notifications = Notification.query.filter(
        Notification.recipient_type == 'all_lecturers'
    ).order_by(Notification.created_at.desc()).limit(10).all()
    
    # Calculate quick stats
    all_sessions = LectureSession.query.join(Unit).filter(
        Unit.lecturer_id == lecturer.id
    ).all()
    
    total_sessions = len(all_sessions)
    active_sessions = len([s for s in all_sessions if s.is_active])
    
    # Get all attendances for lecturer's sessions
    session_ids = [s.id for s in all_sessions]
    total_attendances = Attendance.query.filter(
        Attendance.session_id.in_(session_ids)
    ).count() if session_ids else 0
    
    # This week's stats
    week_start = datetime.utcnow() - timedelta(days=7)
    sessions_this_week = LectureSession.query.join(Unit).filter(
        Unit.lecturer_id == lecturer.id,
        LectureSession.created_at >= week_start
    ).count()
    
    attendances_this_week = Attendance.query.filter(
        Attendance.session_id.in_(session_ids),
        Attendance.submitted_at >= week_start
    ).count() if session_ids else 0
    
    # This month's stats
    month_start = datetime.utcnow() - timedelta(days=30)
    sessions_this_month = LectureSession.query.join(Unit).filter(
        Unit.lecturer_id == lecturer.id,
        LectureSession.created_at >= month_start
    ).count()
    
    attendances_this_month = Attendance.query.filter(
        Attendance.session_id.in_(session_ids),
        Attendance.submitted_at >= month_start
    ).count() if session_ids else 0
    
    # Late arrivals count
    late_count = Attendance.query.filter(
        Attendance.session_id.in_(session_ids),
        Attendance.is_late == True
    ).count() if session_ids else 0
    
    # Calculate attendance rate (average attendance per session)
    attendance_rate = 0
    if total_sessions > 0:
        attendance_rate = round((total_attendances / total_sessions), 1)
    
    # Get attendance data for charts (last 7 days)
    chart_data = []
    for i in range(6, -1, -1):
        date = datetime.utcnow() - timedelta(days=i)
        date_start = date.replace(hour=0, minute=0, second=0, microsecond=0)
        date_end = date.replace(hour=23, minute=59, second=59, microsecond=999999)
        
        day_attendances = Attendance.query.filter(
            Attendance.session_id.in_(session_ids),
            Attendance.submitted_at >= date_start,
            Attendance.submitted_at <= date_end
        ).count() if session_ids else 0
        
        chart_data.append({
            'date': date.strftime('%Y-%m-%d'),
            'day': date.strftime('%a'),
            'count': day_attendances
        })
    
    stats = {
        'total_sessions': total_sessions,
        'active_sessions': active_sessions,
        'total_attendances': total_attendances,
        'sessions_this_week': sessions_this_week,
        'attendances_this_week': attendances_this_week,
        'sessions_this_month': sessions_this_month,
        'attendances_this_month': attendances_this_month,
        'late_count': late_count,
        'attendance_rate': attendance_rate,
        'chart_data': chart_data
    }
    
    return render_template('lecturer_dashboard.html', 
                         lecturer=lecturer, 
                         units=units, 
                         recent_sessions=recent_sessions,
                         notifications=notifications,
                         stats=stats)

@app.route('/lecturer/create-session', methods=['GET', 'POST'])
@lecturer_required
def create_session():
    try:
        lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
        if not lecturer:
            flash('Lecturer profile not found', 'error')
            return redirect(url_for('logout'))
        
        units = Unit.query.filter_by(lecturer_id=lecturer.id).all()
        
        # Check if lecturer has units
        if not units:
            flash('You need to have at least one unit assigned to create a session. Please contact an administrator.', 'error')
            return redirect(url_for('lecturer_dashboard'))
        
        # Handle GET request
        if request.method == 'GET':
            try:
                return render_template('create_session.html', units=units)
            except Exception as e:
                app.logger.error(f"Error rendering create_session template (GET): {str(e)}", exc_info=True)
                flash(f'Error loading page: {str(e)}', 'error')
                return redirect(url_for('lecturer_dashboard'))
        
        # Handle POST request
        if request.method == 'POST':
            try:
                unit_id = request.form.get('unit_id')
                session_name = request.form.get('session_name')
                hall_name = request.form.get('hall_name', '').strip()
                
                # Get and validate coordinates
                try:
                    lat_str = request.form.get('latitude')
                    lng_str = request.form.get('longitude')
                    
                    if not lat_str or not lng_str:
                        flash('Please provide both latitude and longitude coordinates.', 'error')
                        return render_template('create_session.html', units=units)
                    
                    latitude = float(lat_str)
                    longitude = float(lng_str)
                    
                    # Validate coordinate ranges
                    if not (-90 <= latitude <= 90) or not (-180 <= longitude <= 180):
                        flash('Invalid coordinate values. Latitude must be between -90 and 90, longitude between -180 and 180.', 'error')
                        return render_template('create_session.html', units=units)
                    
                    # Check for obviously invalid coordinates (0,0 is in the ocean off Africa, unlikely for a lecture hall)
                    if latitude == 0.0 and longitude == 0.0:
                        flash('Coordinates cannot be (0, 0). Please use your actual location.', 'error')
                        return render_template('create_session.html', units=units)
                        
                except (ValueError, TypeError) as e:
                    flash(f'Invalid latitude or longitude values. Please enter valid numbers. Error: {str(e)}', 'error')
                    return render_template('create_session.html', units=units)
                    
                # Get and validate location radius
                try:
                    location_radius = request.form.get('location_radius', type=int) or 100
                    if location_radius <= 0:
                        raise ValueError("Radius must be greater than 0")
                except (ValueError, TypeError):
                    location_radius = 100  # Default value if invalid
                
                # Get and validate QR expiry minutes
                try:
                    qr_expiry_minutes = min(int(request.form.get('qr_expiry_minutes', 120, type=int)), 1440)
                    if qr_expiry_minutes <= 0:
                        raise ValueError("Expiry duration must be greater than 0")
                except (ValueError, TypeError):
                    qr_expiry_minutes = 120  # Default value if invalid
                
                if not unit_id or not session_name:
                    flash('Please fill in all required fields', 'error')
                    return render_template('create_session.html', units=units)
                
                # Verify unit belongs to lecturer
                unit = Unit.query.get(unit_id)
                if not unit:
                    flash('Invalid unit selected', 'error')
                    return render_template('create_session.html', units=units)
                
                if unit.lecturer_id != lecturer.id:
                    flash('You do not have permission to create sessions for this unit', 'error')
                    return render_template('create_session.html', units=units)
                
                # Generate unique QR code token
                qr_token = secrets.token_urlsafe(32)
                
                # Get base URL (handles both production and local)
                base_url = get_base_url()
                qr_url = f"{base_url}/attendance/{qr_token}"
                
                # Generate QR code image
                try:
                    qr = qrcode.QRCode(version=1, box_size=10, border=5)
                    qr.add_data(qr_url)
                    qr.make(fit=True)
                    img = qr.make_image(fill_color="black", back_color="white")
                    
                    # Convert to base64
                    img_buffer = io.BytesIO()
                    img.save(img_buffer, format='PNG')
                    img_base64 = base64.b64encode(img_buffer.getvalue()).decode()
                except Exception as e:
                    app.logger.error(f"Error generating QR code: {str(e)}", exc_info=True)
                    flash(f'Error generating QR code: {str(e)}', 'error')
                    return render_template('create_session.html', units=units)
                
                # Calculate expiration time using custom duration
                expires_at = datetime.utcnow() + timedelta(minutes=qr_expiry_minutes)
                
                # Get session start time for late tracking (optional)
                start_time_str = request.form.get('session_start_time')
                session_start_time = None
                if start_time_str:
                    try:
                        session_start_time = datetime.strptime(start_time_str, '%Y-%m-%dT%H:%M')
                    except ValueError:
                        pass
                
                late_threshold_str = request.form.get('late_threshold_minutes')
                late_threshold = int(late_threshold_str) if late_threshold_str else 15
                
                # Log coordinates before storing (for debugging)
                app.logger.info(f"Creating session with coordinates - Lat: {latitude}, Lng: {longitude}")
                
                # Create session
                try:
                    session_obj = LectureSession(
                        unit_id=unit_id,
                        session_name=session_name,
                        hall_name=hall_name if hall_name else None,
                        qr_code_token=qr_token,
                        qr_code_data=img_base64,
                        lecture_hall_latitude=latitude,
                        lecture_hall_longitude=longitude,
                        location_radius=location_radius,
                        expires_at=expires_at,
                        session_start_time=session_start_time,
                        late_threshold_minutes=late_threshold
                    )
                    
                    db.session.add(session_obj)
                    db.session.commit()
                    
                    # Verify coordinates were stored correctly
                    app.logger.info(f"Session created - Stored coordinates - Lat: {session_obj.lecture_hall_latitude}, Lng: {session_obj.lecture_hall_longitude}")
                except Exception as e:
                    db.session.rollback()
                    app.logger.error(f"Error creating session in database: {str(e)}", exc_info=True)
                    flash(f'Error saving session to database: {str(e)}. Please try again.', 'error')
                    return render_template('create_session.html', units=units)
                
                # Sync to other databases
                session_data = {
                    'id': session_obj.id,
                    'unit_id': session_obj.unit_id,
                    'session_name': session_obj.session_name,
                    'hall_name': session_obj.hall_name,
                    'qr_code_token': session_obj.qr_code_token,
                    'qr_code_data': session_obj.qr_code_data,
                    'lecture_hall_latitude': session_obj.lecture_hall_latitude,
                    'lecture_hall_longitude': session_obj.lecture_hall_longitude,
                    'location_radius': session_obj.location_radius,
                    'created_at': session_obj.created_at,
                    'expires_at': session_obj.expires_at,
                    'is_active': session_obj.is_active
                }
                # Sync to additional databases (non-blocking)
                try:
                    get_db_sync().sync_session(session_data)
                except Exception as e:
                    app.logger.warning(f"Database sync failed (non-critical): {str(e)}")
                
                flash('QR code generated successfully!', 'success')
                # New session is never expired
                return render_template('view_qr.html', 
                                     session=session_obj, 
                                     qr_url=qr_url,
                                     expiry_minutes=qr_expiry_minutes,
                                     is_expired=False)
            except Exception as e:
                db.session.rollback()
                app.logger.error(f"Unexpected error in create_session (POST): {str(e)}", exc_info=True)
                flash(f'An unexpected error occurred: {str(e)}. Please try again.', 'error')
                return render_template('create_session.html', units=units)
    except Exception as e:
        app.logger.error(f"Unexpected error in create_session route: {str(e)}", exc_info=True)
        flash(f'An error occurred: {str(e)}. Please try again.', 'error')
        return redirect(url_for('lecturer_dashboard'))

@app.route('/attendance/<token>')
def attendance_form(token):
    """Display attendance form for students"""
    try:
        session_obj = LectureSession.query.filter_by(qr_code_token=token).first()
        
        if not session_obj:
            return render_template('error.html', message='Invalid QR code'), 404
        
        if not session_obj.is_active:
            return render_template('error.html', message='This QR code is no longer active'), 403
        
        if datetime.utcnow() > session_obj.expires_at:
            session_obj.is_active = False
            db.session.commit()
            return render_template('error.html', message='This QR code has expired'), 403
        
        # Calculate remaining time and expiration timestamp
        is_expired = datetime.utcnow() > session_obj.expires_at
        if not is_expired:
            remaining = session_obj.expires_at - datetime.utcnow()
            expiry_minutes = int(remaining.total_seconds() / 60)
        else:
            expiry_minutes = 0
        
        expires_timestamp = int(session_obj.expires_at.timestamp() * 1000)  # JavaScript uses milliseconds
        
        return render_template('attendance_form.html', 
                             session=session_obj, 
                             token=token,
                             is_expired=is_expired,
                             expiry_minutes=expiry_minutes,
                             expires_timestamp=expires_timestamp)
    except Exception as e:
        app.logger.error(f"Error in attendance_form: {str(e)}", exc_info=True)
        return render_template('error.html', message=f'An internal error occurred: {str(e)}'), 500

@app.route('/api/submit-attendance', methods=['POST'])
def submit_attendance():
    """Handle attendance submission with geolocation verification"""
    app.logger.info("=" * 60)
    app.logger.info("ATTENDANCE SUBMISSION REQUEST RECEIVED")
    app.logger.info("=" * 60)
    
    data = request.get_json()
    
    if not data:
        app.logger.error("No JSON data received in request")
        return jsonify({'success': False, 'message': 'Invalid request data'}), 400
    
    token = data.get('token')
    admission_no = data.get('admission_no')
    student_name = data.get('student_name')
    latitude = data.get('latitude')
    longitude = data.get('longitude')
    photo_data = data.get('photo')  # Base64 encoded photo
    
    app.logger.info(f"Received data - Token: {token[:20]}..., Admission: {admission_no}, Name: {student_name}")
    app.logger.info(f"Received coordinates - Lat: {latitude}, Lng: {longitude}")
    app.logger.info(f"Photo data received: {'Yes' if photo_data else 'No'}")
    
    if not all([token, admission_no, student_name, latitude, longitude]):
        return jsonify({'success': False, 'message': 'Missing required fields'}), 400
    
    # Photo is now required for attendance submission
    if not photo_data:
        return jsonify({'success': False, 'message': 'Photo is required. Please capture your photo before submitting.'}), 400
    
    # Validate and convert coordinates
    try:
        latitude = float(latitude)
        longitude = float(longitude)
        
        # Add coordinate validation
        if not (-90 <= latitude <= 90) or not (-180 <= longitude <= 180):
            app.logger.error(f"Invalid coordinate range. Lat: {latitude}, Lng: {longitude}")
            return jsonify({'success': False, 'message': 'Invalid coordinate values. Please check your location settings.'}), 400
            
    except (ValueError, TypeError) as e:
        app.logger.error(f"Coordinate conversion error: {str(e)}")
        return jsonify({'success': False, 'message': 'Invalid latitude or longitude values'}), 400
    
    # Find session
    session_obj = LectureSession.query.filter_by(qr_code_token=token).first()
    if not session_obj:
        return jsonify({'success': False, 'message': 'Invalid QR code'}), 404
    
    # Check if expired
    if datetime.utcnow() > session_obj.expires_at:
        session_obj.is_active = False
        db.session.commit()
        return jsonify({'success': False, 'message': 'QR code has expired'}), 403
    
    # Verify location (within lecture hall radius)
    try:
        # Validate coordinates before conversion
        if session_obj.lecture_hall_latitude is None or session_obj.lecture_hall_longitude is None:
            app.logger.error("Lecture hall coordinates are None")
            return jsonify({'success': False, 'message': 'Lecture hall location not set. Please contact your lecturer.'}), 500
        
        if latitude is None or longitude is None:
            app.logger.error("Student coordinates are None")
            return jsonify({'success': False, 'message': 'Invalid location data. Please try again.'}), 400
        
        lecture_lat = float(session_obj.lecture_hall_latitude)
        lecture_lng = float(session_obj.lecture_hall_longitude)
        student_lat = float(latitude)
        student_lng = float(longitude)
        
        # Validate coordinate ranges
        if not (-90 <= lecture_lat <= 90) or not (-180 <= lecture_lng <= 180):
            app.logger.error(f"Invalid lecture hall coordinates: {lecture_lat}, {lecture_lng}")
            return jsonify({'success': False, 'message': 'Invalid lecture hall location. Please contact your lecturer.'}), 500
        
        if not (-90 <= student_lat <= 90) or not (-180 <= student_lng <= 180):
            app.logger.error(f"Invalid student coordinates: {student_lat}, {student_lng}")
            return jsonify({'success': False, 'message': 'Invalid location data. Please try again.'}), 400
        
        # Log coordinates for debugging (with detailed info)
        app.logger.info(f"=== DISTANCE CALCULATION DEBUG ===")
        app.logger.info(f"Lecture hall coordinates - Lat: {lecture_lat} ({type(lecture_lat).__name__}), Lng: {lecture_lng} ({type(lecture_lng).__name__})")
        app.logger.info(f"Student coordinates - Lat: {student_lat} ({type(student_lat).__name__}), Lng: {student_lng} ({type(student_lng).__name__})")
        app.logger.info(f"Session radius: {session_obj.location_radius}m")
        
        # Calculate distance using geodesic (expects lat, lon order)
        lecture_location = (lecture_lat, lecture_lng)
        student_location = (student_lat, student_lng)
        
        distance = geodesic(lecture_location, student_location).meters
        session_radius = float(session_obj.location_radius or app.config['LECTURE_HALL_RADIUS'])
        
        app.logger.info(f"Initial calculated distance: {distance:.2f}m (radius: {session_radius}m)")
        
        # Check if distance is unreasonably large (possible coordinate swap or error)
        # Lower threshold to catch more cases - if distance > 1km, try swapping (lecture halls are typically within 1km)
        if distance > 1000:  # More than 1km is suspicious for a lecture hall
            app.logger.warning(f"Large distance ({distance:.2f}m) detected. Trying all coordinate combinations...")
            
            # Try all possible combinations
            combinations = [
                (geodesic(lecture_location, student_location).meters, 'normal', 'none'),
                (geodesic(lecture_location, (student_lng, student_lat)).meters, 'student_swap', 'student'),
                (geodesic((lecture_lng, lecture_lat), student_location).meters, 'lecture_swap', 'lecture'),
                (geodesic((lecture_lng, lecture_lat), (student_lng, student_lat)).meters, 'both_swap', 'both'),
            ]
            
            # Log all combinations with detailed info
            app.logger.info("All coordinate combinations tested:")
            for dist, name, _ in combinations:
                app.logger.info(f"  {name}: {dist:.2f}m")
            
            # Find the best (smallest) distance - ALWAYS use the smallest one found
            # This ensures we get the best possible result even if all distances are large
            best_distance, best_name, best_swap_type = min(combinations, key=lambda x: x[0])
            
            # ALWAYS use the best distance found (even if it's still large, it's the best we can do)
            # The best_distance is already calculated from the combinations, so use it directly
            if best_distance != distance:
                app.logger.warning(f"Using {best_name} - original: {distance:.2f}m, corrected: {best_distance:.2f}m (improvement: {distance - best_distance:.2f}m)")
            distance = best_distance  # Always use the best distance found
            app.logger.info(f"✓ Using best distance: {best_distance:.2f}m (from {best_name})")
            
            # If best distance is still very large, log a warning
            if best_distance > 50000:  # More than 50km
                app.logger.error(f"⚠ WARNING: Even after coordinate swap detection, distance is still very large: {best_distance:.2f}m")
                app.logger.error(f"   This suggests coordinates may be fundamentally incorrect or in wrong format")
                app.logger.error(f"   Lecture hall: ({lecture_lat}, {lecture_lng})")
                app.logger.error(f"   Student: ({student_lat}, {student_lng})")
        
        app.logger.info(f"Final distance used: {distance:.2f}m (radius: {session_radius}m)")
        app.logger.info(f"=== END DISTANCE CALCULATION ===")
        
        # Add a small buffer (5% or 10m, whichever is larger) to account for GPS accuracy variations
        # This helps when students are right at the edge of the radius
        buffer = max(session_radius * 0.05, 10)  # 5% or 10m buffer
        effective_radius = session_radius + buffer
        
        app.logger.info(f"Effective radius with buffer: {effective_radius:.2f}m (original: {session_radius}m, buffer: {buffer:.2f}m)")
        
        if distance > effective_radius:
            return jsonify({
                'success': False, 
                'message': 'You are too far from the lecture hall. Please move closer to submit attendance.'
            }), 403
    except (ValueError, TypeError) as e:
        app.logger.error(f"Error converting coordinates: {str(e)}")
        return jsonify({'success': False, 'message': 'Invalid coordinate values. Please try again.'}), 400
    except Exception as e:
        app.logger.error(f"Error calculating distance: {str(e)}")
        return jsonify({'success': False, 'message': 'Error verifying location. Please try again.'}), 500
    
    # Check for duplicate submission
    existing = Attendance.query.filter_by(
        session_id=session_obj.id,
        admission_no=admission_no
    ).first()
    
    if existing:
        return jsonify({'success': False, 'message': 'You have already submitted attendance for this session'}), 400
    
    # Get client IP address
    client_ip = request.remote_addr
    # Handle proxy headers (X-Forwarded-For)
    if request.headers.get('X-Forwarded-For'):
        client_ip = request.headers.get('X-Forwarded-For').split(',')[0].strip()
    
    # Device-based rate limiting: Only one submission per device per session
    # Check if this IP/device has already submitted for this session
    existing_device_submission = Attendance.query.filter_by(
        session_id=session_obj.id,
        ip_address=client_ip
    ).first()
    
    if existing_device_submission:
        return jsonify({
            'success': False, 
            'message': 'This device has already been used to submit attendance for this session. Each device can only be used once per session.'
        }), 429  # 429 Too Many Requests
    
    # Additional check: Verify the admission number doesn't already exist for this session
    # (Double-check even though we already checked above)
    duplicate_admission = Attendance.query.filter_by(
        session_id=session_obj.id,
        admission_no=admission_no
    ).first()
    
    if duplicate_admission:
        return jsonify({'success': False, 'message': 'This admission number has already been used for this session'}), 400
    
    # Check for late arrival if session has start time
    is_late = False
    arrival_minutes_late = None
    if session_obj.session_start_time:
        submission_time = datetime.utcnow()
        time_diff = (submission_time - session_obj.session_start_time).total_seconds() / 60
        if time_diff > session_obj.late_threshold_minutes:
            is_late = True
            arrival_minutes_late = int(time_diff)
    
    # Analyze photo for fraud detection
    from photo_analysis import calculate_photo_hash, extract_face_encoding, analyze_photo_for_fraud
    
    photo_hash = calculate_photo_hash(photo_data)
    face_encoding = extract_face_encoding(photo_data)
    
    # Detect matches before saving
    fraud_matches = analyze_photo_for_fraud(photo_data, session_obj.id)
    
    # Create attendance record
    attendance = Attendance(
        session_id=session_obj.id,
        admission_no=admission_no,
        student_name=student_name,
        submission_latitude=latitude,
        submission_longitude=longitude,
        ip_address=get_client_ip(),
        is_late=is_late,
        arrival_minutes_late=arrival_minutes_late,
        photo_data=photo_data,  # Store base64 encoded photo
        photo_hash=photo_hash,
        face_encoding=face_encoding
    )
    
    db.session.add(attendance)
    db.session.flush()  # Get the attendance ID
    
    # Create PhotoMatch records for detected fraud
    from models import PhotoMatch
    
    for duplicate in fraud_matches.get('duplicates', []):
        match = PhotoMatch(
            source_attendance_id=attendance.id,
            target_attendance_id=duplicate['attendance_id'],
            match_type='duplicate',
            similarity_score=None
        )
        db.session.add(match)
    
    for similar_face in fraud_matches.get('similar_faces', []):
        match = PhotoMatch(
            source_attendance_id=attendance.id,
            target_attendance_id=similar_face['attendance_id'],
            match_type='face_similar',
            similarity_score=similar_face['similarity_score']
        )
        db.session.add(match)
    
    db.session.commit()
    
    late_msg = f" ({arrival_minutes_late} minutes late)" if is_late else ""
    log_activity('submit_attendance', 'attendance', attendance.id,
                f'Student {admission_no} submitted attendance for session {session_obj.session_name}{late_msg}')
    
    # Sync to other databases
    attendance_data = {
        'session_id': attendance.session_id,
        'admission_no': attendance.admission_no,
        'student_name': attendance.student_name,
        'submission_latitude': attendance.submission_latitude,
        'submission_longitude': attendance.submission_longitude,
        'submitted_at': attendance.submitted_at,
        'ip_address': attendance.ip_address,
        'photo_data': attendance.photo_data  # Include photo in sync
    }
    # Sync to additional databases (non-blocking)
    try:
        synced_servers = get_db_sync().sync_attendance(attendance_data)
    except Exception as e:
        app.logger.warning(f"Database sync failed (non-critical): {str(e)}")
        synced_servers = []
    
    return jsonify({
        'success': True, 
        'message': 'Attendance submitted successfully!',
        'synced_servers': len(synced_servers)
    })

@app.route('/lecturer/view-attendance/<int:session_id>')
@lecturer_required
def view_attendance(session_id):
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    # Get filter parameters
    search = request.args.get('search', '').strip()
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    status_filter = request.args.get('status', '')  # all, on_time, late
    sort_by = request.args.get('sort_by', 'submitted_at')  # submitted_at, admission_no, student_name
    sort_order = request.args.get('sort_order', 'desc')  # asc, desc
    
    # Build query
    query = Attendance.query.filter_by(session_id=session_id)
    
    # Apply filters
    if start_date:
        start_dt = datetime.strptime(start_date, '%Y-%m-%d')
        query = query.filter(Attendance.submitted_at >= start_dt)
    if end_date:
        end_dt = datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1)
        query = query.filter(Attendance.submitted_at < end_dt)
    if search:
        query = query.filter(
            (Attendance.admission_no.like(f'%{search}%')) |
            (Attendance.student_name.like(f'%{search}%'))
        )
    if status_filter == 'on_time':
        query = query.filter(Attendance.is_late == False)
    elif status_filter == 'late':
        query = query.filter(Attendance.is_late == True)
    
    # Apply sorting
    if sort_by == 'admission_no':
        if sort_order == 'asc':
            query = query.order_by(Attendance.admission_no.asc())
        else:
            query = query.order_by(Attendance.admission_no.desc())
    elif sort_by == 'student_name':
        if sort_order == 'asc':
            query = query.order_by(Attendance.student_name.asc())
        else:
            query = query.order_by(Attendance.student_name.desc())
    else:  # submitted_at
        if sort_order == 'asc':
            query = query.order_by(Attendance.submitted_at.asc())
        else:
            query = query.order_by(Attendance.submitted_at.desc())
    
    attendances = query.all()
    
    # Get fraud detection matches for this session
    from models import PhotoMatch
    attendance_ids = [a.id for a in attendances]
    fraud_matches = PhotoMatch.query.filter(
        (PhotoMatch.source_attendance_id.in_(attendance_ids)) |
        (PhotoMatch.target_attendance_id.in_(attendance_ids))
    ).all()
    
    # Create a map of attendance_id -> list of matches
    fraud_map = {}
    for match in fraud_matches:
        if match.source_attendance_id not in fraud_map:
            fraud_map[match.source_attendance_id] = []
        if match.target_attendance_id not in fraud_map:
            fraud_map[match.target_attendance_id] = []
        fraud_map[match.source_attendance_id].append(match)
        fraud_map[match.target_attendance_id].append(match)
    
    # Calculate statistics
    total_count = len(attendances)
    on_time_count = len([a for a in attendances if not a.is_late])
    late_count = len([a for a in attendances if a.is_late])
    
    # Count unverified matches
    unverified_count = PhotoMatch.query.filter(
        (PhotoMatch.source_attendance_id.in_(attendance_ids)) |
        (PhotoMatch.target_attendance_id.in_(attendance_ids))
    ).filter_by(verified_by_lecturer=None).count()
    
    return render_template('view_attendance.html', 
                         session=session_obj, 
                         attendances=attendances,
                         search=search,
                         start_date=start_date,
                         end_date=end_date,
                         status_filter=status_filter,
                         sort_by=sort_by,
                         sort_order=sort_order,
                         total_count=total_count,
                         on_time_count=on_time_count,
                         late_count=late_count,
                         fraud_map=fraud_map,
                         unverified_matches_count=unverified_count)

@app.route('/lecturer/export-attendance/<int:session_id>')
@lecturer_required
def export_attendance(session_id):
    """Export attendance records to Excel format"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    # Get all attendance records
    attendances = Attendance.query.filter_by(session_id=session_id).order_by(
        Attendance.submitted_at.asc()
    ).all()
    
    # Create a new workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    
    # Define header style
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Add session information
    ws['A1'] = 'Session Information'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:E1')
    
    ws['A2'] = 'Session Name:'
    ws['B2'] = session_obj.session_name
    ws['A3'] = 'Unit:'
    ws['B3'] = f"{session_obj.unit.unit_code} - {session_obj.unit.unit_name}"
    ws['A4'] = 'Created:'
    ws['B4'] = session_obj.created_at.strftime('%Y-%m-%d %H:%M:%S')
    ws['A5'] = 'Expires:'
    ws['B5'] = session_obj.expires_at.strftime('%Y-%m-%d %H:%M:%S')
    ws['A6'] = 'Status:'
    ws['B6'] = 'Active' if session_obj.is_active else 'Expired'
    ws['A7'] = 'Total Students:'
    ws['B7'] = len(attendances)
    
    # Add empty row
    ws['A9'] = ''
    
    # Add table headers
    headers = ['#', 'Admission Number', 'Student Name', 'Submitted At', 'Status', 'Location (Lat, Long)', 'IP Address']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=10, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Add attendance data
    # Track IP addresses for duplicate detection
    ip_counts = {}
    for attendance in attendances:
        ip = attendance.ip_address or 'N/A'
        ip_counts[ip] = ip_counts.get(ip, 0) + 1
    
    for idx, attendance in enumerate(attendances, 1):
        row = 10 + idx
        ip = attendance.ip_address or 'N/A'
        ip_count = ip_counts.get(ip, 0)
        
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=attendance.admission_no)
        ws.cell(row=row, column=3, value=attendance.student_name)
        ws.cell(row=row, column=4, value=attendance.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
        
        # Status column (Late/On Time)
        status_cell = ws.cell(row=row, column=5)
        if attendance.is_late:
            status_cell.value = f"Late ({attendance.arrival_minutes_late}m)"
            status_cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
            status_cell.font = Font(color="FFFFFF", bold=True)
        else:
            status_cell.value = "On Time"
            status_cell.fill = PatternFill(start_color="51CF66", end_color="51CF66", fill_type="solid")
            status_cell.font = Font(color="FFFFFF", bold=True)
        
        ws.cell(row=row, column=6, value=f"{attendance.submission_latitude:.6f}, {attendance.submission_longitude:.6f}")
        ip_cell = ws.cell(row=row, column=7, value=ip)
        if ip_count > 1:
            ip_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            ip_cell.value = f"{ip} ({ip_count}x)"
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 15
    
    # Center align the number column
    for row in range(11, 11 + len(attendances)):
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
    
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Create filename with session name and date
    safe_session_name = "".join(c for c in session_obj.session_name if c.isalnum() or c in (' ', '-', '_')).rstrip()
    filename = f"Attendance_{safe_session_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@app.route('/admin/manage-lecturers', methods=['GET', 'POST'])
@admin_required
def manage_lecturers():
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'create':
            username = request.form.get('username')
            email = request.form.get('email')
            password = request.form.get('password')
            full_name = request.form.get('full_name')
            employee_id = request.form.get('employee_id')
            
            if User.query.filter_by(username=username).first():
                flash('Username already exists', 'error')
                return redirect(url_for('manage_lecturers'))
            
            user = User(username=username, email=email, role='lecturer')
            user.set_password(password)
            db.session.add(user)
            db.session.flush()
            
            department_id = request.form.get('department_id')
            department_id = int(department_id) if department_id else None
            
            lecturer = Lecturer(
                user_id=user.id,
                full_name=full_name,
                employee_id=employee_id,
                department_id=department_id
            )
            db.session.add(lecturer)
            db.session.commit()
            log_activity('create_lecturer', 'lecturer', lecturer.id, 
                        f'Created lecturer {full_name} ({employee_id})')
            flash('Lecturer created successfully', 'success')
        
        elif action == 'bulk_assign_department':
            # Bulk assign lecturers to department
            lecturer_ids = request.form.getlist('lecturer_ids')
            department_id = request.form.get('bulk_department_id')
            department_id = int(department_id) if department_id else None
            
            count = 0
            for lecturer_id in lecturer_ids:
                lecturer = Lecturer.query.get(lecturer_id)
                if lecturer:
                    lecturer.department_id = department_id
                    count += 1
            
            db.session.commit()
            log_activity('bulk_assign_department', description=f'Bulk assigned {count} lecturer(s) to department')
            flash(f'Successfully assigned {count} lecturer(s) to department', 'success')
        
        elif action == 'bulk_delete_lecturers':
            lecturer_ids = request.form.getlist('lecturer_ids')
            count = 0
            for lecturer_id in lecturer_ids:
                lecturer = Lecturer.query.get(lecturer_id)
                if lecturer:
                    user = lecturer.user
                    db.session.delete(lecturer)
                    db.session.delete(user)
                    count += 1
            db.session.commit()
            log_activity('bulk_delete_lecturers', description=f'Bulk deleted {count} lecturer(s)')
            flash(f'Successfully deleted {count} lecturer(s)', 'success')
        
        elif action == 'edit':
            lecturer_id = request.form.get('lecturer_id')
            lecturer = Lecturer.query.get_or_404(lecturer_id)
            user = lecturer.user
            
            # Update lecturer info
            lecturer.full_name = request.form.get('full_name')
            lecturer.employee_id = request.form.get('employee_id')
            
            # Update department
            department_id = request.form.get('department_id')
            lecturer.department_id = int(department_id) if department_id else None
            
            # Update user info
            new_username = request.form.get('username')
            new_email = request.form.get('email')
            
            # Check if username is being changed and if it's already taken
            if new_username != user.username:
                if User.query.filter_by(username=new_username).first():
                    flash('Username already exists', 'error')
                    return redirect(url_for('manage_lecturers'))
                user.username = new_username
            
            # Check if email is being changed and if it's already taken
            if new_email != user.email:
                if User.query.filter_by(email=new_email).first():
                    flash('Email already exists', 'error')
                    return redirect(url_for('manage_lecturers'))
                user.email = new_email
            
            # Update password if provided
            new_password = request.form.get('password')
            if new_password and new_password.strip():
                user.set_password(new_password)
            
            db.session.commit()
            flash('Lecturer updated successfully', 'success')
        
        elif action == 'delete':
            lecturer_id = request.form.get('lecturer_id')
            lecturer = Lecturer.query.get_or_404(lecturer_id)
            user = lecturer.user
            db.session.delete(lecturer)
            db.session.delete(user)
            db.session.commit()
            flash('Lecturer deleted successfully', 'success')
        
        # Department actions
        elif action == 'create_department':
            name = request.form.get('name')
            code = request.form.get('code')
            description = request.form.get('description', '').strip()
            
            if Department.query.filter_by(name=name).first():
                flash('Department name already exists', 'error')
                return redirect(url_for('manage_lecturers'))
            
            if Department.query.filter_by(code=code).first():
                flash('Department code already exists', 'error')
                return redirect(url_for('manage_lecturers'))
            
            department = Department(
                name=name,
                code=code.upper(),
                description=description
            )
            db.session.add(department)
            db.session.commit()
            flash('Department created successfully', 'success')
        
        elif action == 'edit_department':
            department_id = request.form.get('department_id')
            department = Department.query.get_or_404(department_id)
            
            old_name = department.name
            old_code = department.code
            
            new_name = request.form.get('name')
            new_code = request.form.get('code')
            new_description = request.form.get('description', '').strip()
            
            # Check if department name is being changed and if it's already taken
            if new_name != old_name:
                if Department.query.filter_by(name=new_name).first():
                    flash('Department name already exists', 'error')
                    return redirect(url_for('manage_lecturers'))
                department.name = new_name
            
            # Check if department code is being changed and if it's already taken
            if new_code.upper() != old_code:
                if Department.query.filter_by(code=new_code.upper()).first():
                    flash('Department code already exists', 'error')
                    return redirect(url_for('manage_lecturers'))
                department.code = new_code.upper()
            
            department.description = new_description
            
            db.session.commit()
            log_activity('edit_department', 'department', department.id,
                        f'Updated department {old_code} -> {new_code}')
            flash('Department updated successfully', 'success')
        
        elif action == 'delete_department':
            department_id = request.form.get('department_id')
            department = Department.query.get_or_404(department_id)
            
            if department.lecturers:
                flash(f'Cannot delete department. {len(department.lecturers)} lecturer(s) are assigned to it.', 'error')
                return redirect(url_for('manage_lecturers'))
            
            dept_code = department.code
            db.session.delete(department)
            db.session.commit()
            log_activity('delete_department', 'department', department_id,
                        f'Deleted department {dept_code}')
            flash('Department deleted successfully', 'success')
        
        # Unit actions
        elif action == 'create_unit':
            unit_code = request.form.get('unit_code')
            unit_name = request.form.get('unit_name')
            lecturer_id = request.form.get('lecturer_id')
            
            if Unit.query.filter_by(unit_code=unit_code).first():
                flash('Unit code already exists', 'error')
                return redirect(url_for('manage_lecturers'))
            
            unit = Unit(
                unit_code=unit_code,
                unit_name=unit_name,
                lecturer_id=lecturer_id
            )
            db.session.add(unit)
            db.session.commit()
            flash('Unit created successfully', 'success')
        
        elif action == 'edit_unit':
            unit_id = request.form.get('unit_id')
            unit = Unit.query.get_or_404(unit_id)
            
            old_code = unit.unit_code
            old_name = unit.unit_name
            
            new_code = request.form.get('unit_code')
            new_name = request.form.get('unit_name')
            new_lecturer_id = request.form.get('lecturer_id')
            
            # Check if unit code is being changed and if it's already taken
            if new_code != old_code:
                if Unit.query.filter_by(unit_code=new_code).first():
                    flash('Unit code already exists', 'error')
                    return redirect(url_for('manage_lecturers'))
                unit.unit_code = new_code
            
            unit.unit_name = new_name
            unit.lecturer_id = int(new_lecturer_id) if new_lecturer_id else None
            
            db.session.commit()
            log_activity('edit_unit', 'unit', unit.id, 
                        f'Updated unit {old_code} -> {new_code}')
            flash('Unit updated successfully', 'success')
        
        elif action == 'create_unit':
            unit_code = request.form.get('unit_code')
            unit_name = request.form.get('unit_name')
            lecturer_id = request.form.get('lecturer_id')
            
            if Unit.query.filter_by(unit_code=unit_code).first():
                flash('Unit code already exists', 'error')
                return redirect(url_for('manage_lecturers'))
            
            unit = Unit(
                unit_code=unit_code,
                unit_name=unit_name,
                lecturer_id=lecturer_id
            )
            db.session.add(unit)
            db.session.commit()
            log_activity('create_unit', 'unit', unit.id, 
                        f'Created unit {unit_code}')
            flash('Unit created successfully', 'success')
        
        elif action == 'delete_unit':
            unit_id = request.form.get('unit_id')
            unit = Unit.query.get_or_404(unit_id)
            unit_code = unit.unit_code
            db.session.delete(unit)
            db.session.commit()
            log_activity('delete_unit', 'unit', unit_id, 
                        f'Deleted unit {unit_code}')
            flash('Unit deleted successfully', 'success')
    
    lecturers = Lecturer.query.all()
    departments = Department.query.order_by(Department.name).all()
    units = Unit.query.all()
    return render_template('manage_lecturers.html', 
                         lecturers=lecturers, 
                         departments=departments,
                         units=units)

@app.route('/admin/manage-departments', methods=['GET', 'POST'])
@admin_required
def manage_departments():
    """Manage departments"""
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'create':
            name = request.form.get('name')
            code = request.form.get('code')
            description = request.form.get('description', '').strip()
            
            if Department.query.filter_by(name=name).first():
                flash('Department name already exists', 'error')
                return redirect(url_for('manage_departments'))
            
            if Department.query.filter_by(code=code).first():
                flash('Department code already exists', 'error')
                return redirect(url_for('manage_departments'))
            
            department = Department(
                name=name,
                code=code.upper(),
                description=description
            )
            db.session.add(department)
            db.session.commit()
            
            flash('Department created successfully', 'success')
        
        elif action == 'delete':
            department_id = request.form.get('department_id')
            department = Department.query.get_or_404(department_id)
            
            # Check if department has lecturers
            if department.lecturers:
                flash(f'Cannot delete department. {len(department.lecturers)} lecturer(s) are assigned to it.', 'error')
                return redirect(url_for('manage_departments'))
            
            db.session.delete(department)
            db.session.commit()
            flash('Department deleted successfully', 'success')
    
    departments = Department.query.order_by(Department.name).all()
    return render_template('manage_departments.html', departments=departments)

@app.route('/admin/manage-units', methods=['GET', 'POST'])
@admin_required
def manage_units():
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'create':
            unit_code = request.form.get('unit_code')
            unit_name = request.form.get('unit_name')
            lecturer_id = request.form.get('lecturer_id')
            
            if Unit.query.filter_by(unit_code=unit_code).first():
                flash('Unit code already exists', 'error')
                return redirect(url_for('manage_units'))
            
            unit = Unit(
                unit_code=unit_code,
                unit_name=unit_name,
                lecturer_id=lecturer_id
            )
            db.session.add(unit)
            db.session.commit()
            flash('Unit created successfully', 'success')
        
        elif action == 'delete':
            unit_id = request.form.get('unit_id')
            unit = Unit.query.get_or_404(unit_id)
            db.session.delete(unit)
            db.session.commit()
            flash('Unit deleted successfully', 'success')
    
    units = Unit.query.all()
    lecturers = Lecturer.query.all()
    return render_template('manage_units.html', units=units, lecturers=lecturers)

@app.route('/lecturer/sessions')
@lecturer_required
def lecturer_sessions():
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    sessions = LectureSession.query.join(Unit).filter(
        Unit.lecturer_id == lecturer.id
    ).order_by(LectureSession.created_at.desc()).all()
    
    return render_template('lecturer_sessions.html', sessions=sessions)

@app.route('/lecturer/view-qr/<int:session_id>')
@lecturer_required
def view_session_qr(session_id):
    """View QR code for an existing session - only the owner lecturer can view it"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit (security check)
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied. You can only view QR codes for your own sessions.', 'error')
        return redirect(url_for('lecturer_sessions'))
    
    # Check if session has expired
    is_expired = datetime.utcnow() > session_obj.expires_at
    if is_expired:
        # Update session status if expired
        if session_obj.is_active:
            session_obj.is_active = False
            db.session.commit()
    
    # Reconstruct QR URL
    base_url = get_base_url()
    qr_url = f"{base_url}/attendance/{session_obj.qr_code_token}"
    
    # Calculate remaining time until expiration
    if not is_expired:
        remaining = session_obj.expires_at - datetime.utcnow()
        expiry_minutes = int(remaining.total_seconds() / 60)
    else:
        expiry_minutes = 0
    
    # Convert expires_at to timestamp for JavaScript
    expires_timestamp = int(session_obj.expires_at.timestamp() * 1000)  # JavaScript uses milliseconds
    
    return render_template('view_qr.html', 
                         session=session_obj, 
                         qr_url=qr_url,
                         expiry_minutes=expiry_minutes,
                         is_expired=is_expired,
                         expires_timestamp=expires_timestamp)

@app.route('/lecturer/duplicate-session/<int:session_id>')
@lecturer_required
def duplicate_session(session_id):
    """Duplicate an existing session with new settings"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    original_session = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if original_session.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    # Create new session with same settings
    new_session = LectureSession(
        session_name=f"{original_session.session_name} (Copy)",
        unit_id=original_session.unit_id,
        lecture_hall_latitude=original_session.lecture_hall_latitude,
        lecture_hall_longitude=original_session.lecture_hall_longitude,
        location_radius=original_session.location_radius,
        hall_name=original_session.hall_name,
        qr_code_token=secrets.token_urlsafe(32),
        is_active=False,  # Start as inactive
        expires_at=datetime.utcnow() + timedelta(minutes=app.config['QR_CODE_EXPIRY_MINUTES']),
        session_start_time=original_session.session_start_time,
        late_threshold_minutes=original_session.late_threshold_minutes
    )
    
    db.session.add(new_session)
    db.session.commit()
    
    flash(f'Session duplicated successfully! New session: {new_session.session_name}', 'success')
    return redirect(url_for('lecturer_sessions'))

@app.route('/lecturer/extend-session/<int:session_id>', methods=['GET', 'POST'])
@lecturer_required
def extend_session(session_id):
    """Extend session expiration time"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    if request.method == 'POST':
        minutes = request.form.get('minutes', type=int)
        if minutes and minutes > 0:
            session_obj.expires_at = datetime.utcnow() + timedelta(minutes=minutes)
            session_obj.is_active = True
            db.session.commit()
            # Format expiration time for display
            local_tz = timezone(timedelta(hours=TIMEZONE_OFFSET_HOURS))
            local_expiry = session_obj.expires_at.replace(tzinfo=timezone.utc).astimezone(local_tz)
            flash(f'Session extended by {minutes} minutes. New expiration: {local_expiry.strftime("%Y-%m-%d %H:%M:%S")}', 'success')
            return redirect(url_for('view_session_qr', session_id=session_id))
        else:
            flash('Invalid minutes value', 'error')
    
    return render_template('extend_session.html', session=session_obj)

@app.route('/lecturer/delete-session/<int:session_id>', methods=['POST'])
@lecturer_required
def delete_session(session_id):
    """Delete a lecture session - only the owner lecturer can delete it"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit (security check)
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied. You can only delete your own sessions.', 'error')
        return redirect(url_for('lecturer_sessions'))
    
    # Delete the session (cascade will delete all attendances)
    session_name = session_obj.session_name
    db.session.delete(session_obj)
    db.session.commit()
    
    # Sync deletion to other databases
    try:
        get_db_sync().delete_session(session_id)
    except Exception as e:
        # Log error but don't fail the deletion
        app.logger.error(f"Failed to sync session deletion: {str(e)}")
    
    flash(f'Session "{session_name}" deleted successfully', 'success')
    return redirect(url_for('lecturer_sessions'))

@app.route('/admin/send-notification', methods=['GET', 'POST'])
@admin_required
def send_notification():
    """Send notification to all lecturers"""
    if request.method == 'POST':
        title = request.form.get('title')
        message = request.form.get('message')
        
        if not title or not message:
            flash('Please fill in both title and message', 'error')
            return render_template('send_notification.html')
        
        # Create notification
        notification = Notification(
            title=title.strip(),
            message=message.strip(),
            sender_id=current_user.id,
            recipient_type='all_lecturers'
        )
        
        db.session.add(notification)
        db.session.commit()
        
        # Get count of lecturers
        lecturer_count = Lecturer.query.count()
        
        flash(f'Notification sent successfully to {lecturer_count} lecturer(s)!', 'success')
        return redirect(url_for('admin_dashboard'))
    
    return render_template('send_notification.html')

@app.route('/lecturer/notifications')
@lecturer_required
def view_notifications():
    """View all notifications for lecturer"""
    notifications = Notification.query.filter(
        Notification.recipient_type == 'all_lecturers'
    ).order_by(Notification.created_at.desc()).all()
    
    return render_template('lecturer_notifications.html', notifications=notifications)

@app.route('/lecturer/notification/<int:notification_id>')
@lecturer_required
def view_notification(notification_id):
    """View a specific notification"""
    notification = Notification.query.get_or_404(notification_id)
    
    # Verify it's a lecturer notification
    if notification.recipient_type != 'all_lecturers':
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    return render_template('view_notification.html', notification=notification)

@app.route('/admin/activity-logs')
@admin_required
def activity_logs():
    """View activity logs"""
    page = request.args.get('page', 1, type=int)
    per_page = 50
    
    logs = ActivityLog.query.order_by(ActivityLog.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    return render_template('activity_logs.html', logs=logs)

@app.route('/admin/login-history')
@admin_required
def login_history():
    """View login history"""
    page = request.args.get('page', 1, type=int)
    per_page = 50
    
    history = LoginHistory.query.order_by(LoginHistory.login_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    return render_template('login_history.html', history=history)

@app.route('/admin/student-attendance-history')
@admin_required
def student_attendance_history():
    """View attendance history for a specific student"""
    admission_no = request.args.get('admission_no', '').strip()
    page = request.args.get('page', 1, type=int)
    per_page = 50
    
    query = Attendance.query
    
    if admission_no:
        query = query.filter(Attendance.admission_no == admission_no)
    
    attendances = query.order_by(Attendance.submitted_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # Get statistics for the student
    stats = None
    if admission_no:
        total_attendance = Attendance.query.filter_by(admission_no=admission_no).count()
        total_sessions = LectureSession.query.count()
        late_count = Attendance.query.filter_by(admission_no=admission_no, is_late=True).count()
        
        # Get unique sessions attended
        sessions_attended = db.session.query(LectureSession).join(Attendance).filter(
            Attendance.admission_no == admission_no
        ).distinct().count()
        
        attendance_percentage = round((sessions_attended / total_sessions * 100), 1) if total_sessions > 0 else 0
        
        stats = {
            'total_attendance': total_attendance,
            'sessions_attended': sessions_attended,
            'total_sessions': total_sessions,
            'attendance_percentage': attendance_percentage,
            'late_count': late_count
        }
    
    return render_template('student_attendance_history.html', 
                         attendances=attendances, 
                         admission_no=admission_no,
                         stats=stats)

@app.route('/lecturer/review-matches/<int:session_id>')
@lecturer_required
def review_photo_matches(session_id):
    """Review detected photo matches for fraud verification"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    from models import PhotoMatch
    # Get all matches for this session
    attendance_ids = [a.id for a in session_obj.attendances]
    
    # Get unverified matches
    unverified_matches = PhotoMatch.query.filter(
        ((PhotoMatch.source_attendance_id.in_(attendance_ids)) |
         (PhotoMatch.target_attendance_id.in_(attendance_ids)))
    ).filter_by(verified_by_lecturer=None).order_by(PhotoMatch.detected_at.desc()).all()
    
    # Get verified matches
    verified_matches = PhotoMatch.query.filter(
        ((PhotoMatch.source_attendance_id.in_(attendance_ids)) |
         (PhotoMatch.target_attendance_id.in_(attendance_ids)))
    ).filter(PhotoMatch.verified_by_lecturer.isnot(None)).order_by(PhotoMatch.verified_at.desc()).all()
    
    return render_template('review_matches.html',
                         session=session_obj,
                         unverified_matches=unverified_matches,
                         verified_matches=verified_matches)

@app.route('/lecturer/verify-match/<int:match_id>', methods=['POST'])
@lecturer_required
def verify_photo_match(match_id):
    """Verify or dismiss a photo match"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        return jsonify({'success': False, 'message': 'Lecturer profile not found'}), 400
    
    from models import PhotoMatch
    match = PhotoMatch.query.get_or_404(match_id)
    
    # Verify lecturer has access to this match
    source_session = match.source_attendance.session
    target_session = match.target_attendance.session
    
    if source_session.unit.lecturer_id != lecturer.id or target_session.unit.lecturer_id != lecturer.id:
        return jsonify({'success': False, 'message': 'Access denied'}), 403
    
    action = request.form.get('action')  # 'confirm' or 'dismiss'
    notes = request.form.get('notes', '').strip()
    
    if action == 'confirm':
        match.verified_by_lecturer = True
    elif action == 'dismiss':
        match.verified_by_lecturer = False
    else:
        return jsonify({'success': False, 'message': 'Invalid action'}), 400
    
    match.verified_at = datetime.utcnow()
    match.verified_by_user_id = current_user.id
    match.notes = notes
    
    db.session.commit()
    
    log_activity('verify_photo_match', 'photo_match', match.id,
                f'Match {match_id} verified as {action} by lecturer')
    
    return jsonify({'success': True, 'message': f'Match {action}ed successfully'})

@app.route('/lecturer/export-attendance-csv/<int:session_id>')
@lecturer_required
def export_attendance_csv(session_id):
    """Export attendance records to CSV format"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    # Get all attendance records
    attendances = Attendance.query.filter_by(session_id=session_id).order_by(
        Attendance.submitted_at.asc()
    ).all()
    
    # Create CSV content
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Write header
    writer.writerow(['#', 'Admission Number', 'Student Name', 'Submitted At', 'Status', 
                     'Latitude', 'Longitude', 'IP Address'])
    
    # Write data
    for idx, attendance in enumerate(attendances, 1):
        status = f"Late ({attendance.arrival_minutes_late}m)" if attendance.is_late else "On Time"
        writer.writerow([
            idx,
            attendance.admission_no,
            attendance.student_name,
            attendance.submitted_at.strftime('%Y-%m-%d %H:%M:%S'),
            status,
            f"{attendance.submission_latitude:.6f}",
            f"{attendance.submission_longitude:.6f}",
            attendance.ip_address or 'N/A'
        ])
    
    output.seek(0)
    
    # Generate filename
    filename = f"attendance_{session_obj.session_name.replace(' ', '_')}_{session_id}.csv"
    
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=filename
    )

@app.route('/lecturer/export-attendance-pdf/<int:session_id>')
@lecturer_required
def export_attendance_pdf(session_id):
    """Export attendance records to PDF format"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    session_obj = LectureSession.query.get_or_404(session_id)
    
    # Verify lecturer owns this unit
    if session_obj.unit.lecturer_id != lecturer.id:
        flash('Access denied', 'error')
        return redirect(url_for('lecturer_dashboard'))
    
    # Get all attendance records
    attendances = Attendance.query.filter_by(session_id=session_id).order_by(
        Attendance.submitted_at.asc()
    ).all()
    
    # Create PDF in memory
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    elements = []
    
    # Define styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('#0066cc'),
        spaceAfter=30,
        alignment=1  # Center
    )
    
    # Add title
    title = Paragraph(f"Attendance Report: {session_obj.session_name}", title_style)
    elements.append(title)
    elements.append(Spacer(1, 0.2*inch))
    
    # Add session information
    info_data = [
        ['Session Name:', session_obj.session_name],
        ['Unit:', f"{session_obj.unit.unit_code} - {session_obj.unit.unit_name}"],
        ['Created:', session_obj.created_at.strftime('%Y-%m-%d %H:%M:%S')],
        ['Expires:', session_obj.expires_at.strftime('%Y-%m-%d %H:%M:%S')],
        ['Status:', 'Active' if session_obj.is_active else 'Expired'],
        ['Total Students:', str(len(attendances))]
    ]
    
    info_table = Table(info_data, colWidths=[2*inch, 4*inch])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ('BACKGROUND', (1, 0), (1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Add attendance table
    if attendances:
        table_data = [['#', 'Admission No', 'Student Name', 'Submitted At', 'Status']]
        
        for idx, attendance in enumerate(attendances, 1):
            status = f"Late ({attendance.arrival_minutes_late}m)" if attendance.is_late else "On Time"
            table_data.append([
                str(idx),
                attendance.admission_no,
                attendance.student_name,
                attendance.submitted_at.strftime('%Y-%m-%d %H:%M'),
                status
            ])
        
        attendance_table = Table(table_data, colWidths=[0.5*inch, 1.5*inch, 2.5*inch, 1.5*inch, 1*inch])
        attendance_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey])
        ]))
        elements.append(attendance_table)
    else:
        elements.append(Paragraph("No attendance records found.", styles['Normal']))
    
    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    
    filename = f"attendance_{session_obj.session_name.replace(' ', '_')}_{session_id}.pdf"
    
    return send_file(
        buffer,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=filename
    )

@app.route('/lecturer/student-attendance-history')
@lecturer_required
def lecturer_student_attendance_history():
    """View attendance history for a specific student (Lecturer view)"""
    lecturer = Lecturer.query.filter_by(user_id=current_user.id).first()
    if not lecturer:
        flash('Lecturer profile not found', 'error')
        return redirect(url_for('logout'))
    
    admission_no = request.args.get('admission_no', '').strip()
    unit_id = request.args.get('unit_id', type=int)
    page = request.args.get('page', 1, type=int)
    per_page = 50
    
    # Get lecturer's units
    lecturer_units = Unit.query.filter_by(lecturer_id=lecturer.id).all()
    unit_ids = [u.id for u in lecturer_units]
    
    # Build query - only for lecturer's sessions
    query = Attendance.query.join(LectureSession).join(Unit).filter(
        Unit.id.in_(unit_ids)
    )
    
    if admission_no:
        query = query.filter(Attendance.admission_no == admission_no)
    
    if unit_id and unit_id in unit_ids:
        query = query.filter(Unit.id == unit_id)
    
    attendances = query.order_by(Attendance.submitted_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # Get statistics for the student
    stats = None
    if admission_no:
        # Only count attendances from lecturer's sessions
        total_attendance = Attendance.query.join(LectureSession).join(Unit).filter(
            Unit.id.in_(unit_ids),
            Attendance.admission_no == admission_no
        ).count()
        
        total_sessions = LectureSession.query.join(Unit).filter(
            Unit.id.in_(unit_ids)
        ).count()
        
        late_count = Attendance.query.join(LectureSession).join(Unit).filter(
            Unit.id.in_(unit_ids),
            Attendance.admission_no == admission_no,
            Attendance.is_late == True
        ).count()
        
        # Get unique sessions attended
        sessions_attended = db.session.query(LectureSession).join(Attendance).join(Unit).filter(
            Unit.id.in_(unit_ids),
            Attendance.admission_no == admission_no
        ).distinct().count()
        
        attendance_percentage = round((sessions_attended / total_sessions * 100), 1) if total_sessions > 0 else 0
        
        stats = {
            'total_attendance': total_attendance,
            'sessions_attended': sessions_attended,
            'total_sessions': total_sessions,
            'attendance_percentage': attendance_percentage,
            'late_count': late_count
        }
    
    return render_template('lecturer_student_history.html', 
                         attendances=attendances, 
                         admission_no=admission_no,
                         unit_id=unit_id,
                         lecturer_units=lecturer_units,
                         stats=stats)

@app.route('/admin/export-attendance-by-department')
@admin_required
def export_attendance_by_department():
    """Export attendance grouped by department"""
    department_id = request.args.get('department_id', type=int)
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    
    # Build query
    query = Attendance.query.join(LectureSession).join(Unit).join(Lecturer)
    
    if department_id:
        query = query.filter(Lecturer.department_id == department_id)
    
    if start_date:
        query = query.filter(Attendance.submitted_at >= datetime.strptime(start_date, '%Y-%m-%d'))
    if end_date:
        query = query.filter(Attendance.submitted_at < datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1))
    
    attendances = query.order_by(Attendance.submitted_at.desc()).all()
    
    # Group by department
    dept_attendance = defaultdict(list)
    
    for attendance in attendances:
        dept = attendance.session.unit.lecturer.department
        dept_name = dept.name if dept else 'No Department'
        dept_attendance[dept_name].append(attendance)
    
    # Create Excel workbook
    wb = Workbook()
    
    # Create summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    
    ws_summary['A1'] = 'Department'
    ws_summary['B1'] = 'Total Attendance'
    ws_summary['C1'] = 'Unique Students'
    ws_summary['D1'] = 'Late Arrivals'
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col in ['A', 'B', 'C', 'D']:
        cell = ws_summary[f'{col}1']
        cell.fill = header_fill
        cell.font = header_font
    
    row = 2
    for dept_name, dept_attendances in sorted(dept_attendance.items()):
        unique_students = len(set(a.admission_no for a in dept_attendances))
        late_count = sum(1 for a in dept_attendances if a.is_late)
        
        ws_summary[f'A{row}'] = dept_name
        ws_summary[f'B{row}'] = len(dept_attendances)
        ws_summary[f'C{row}'] = unique_students
        ws_summary[f'D{row}'] = late_count
        row += 1
    
    # Create sheet for each department
    for dept_name, dept_attendances in sorted(dept_attendance.items()):
        ws = wb.create_sheet(title=dept_name[:31])  # Excel sheet name limit
        
        # Headers
        headers = ['#', 'Admission No', 'Student Name', 'Session', 'Unit', 'Submitted At', 'Status', 'IP Address']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
        
        # Data
        for idx, attendance in enumerate(dept_attendances, 1):
            status = f"Late ({attendance.arrival_minutes_late}m)" if attendance.is_late else "On Time"
            ws.cell(row=idx+1, column=1, value=idx)
            ws.cell(row=idx+1, column=2, value=attendance.admission_no)
            ws.cell(row=idx+1, column=3, value=attendance.student_name)
            ws.cell(row=idx+1, column=4, value=attendance.session.session_name)
            ws.cell(row=idx+1, column=5, value=attendance.session.unit.unit_code)
            ws.cell(row=idx+1, column=6, value=attendance.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=idx+1, column=7, value=status)
            ws.cell(row=idx+1, column=8, value=attendance.ip_address or 'N/A')
    
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"attendance_by_department_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

if __name__ == '__main__':
    with app.app_context():
        try:
            print("Initializing database...")
            db.create_all()
            print("✓ Database tables created/verified")
            
            # Create default admin user if doesn't exist
            try:
                if not User.query.filter_by(username='admin').first():
                    admin = User(username='admin', email='admin@university.edu', role='admin')
                    admin.set_password('admin123')  # Change this in production!
                    db.session.add(admin)
                    db.session.commit()
                    print("✓ Default admin user created: username='admin', password='admin123'")
                else:
                    print("✓ Admin user already exists")
            except Exception as e:
                print(f"⚠ Warning: Could not create admin user: {str(e)}")
                print("You may need to create it manually or check database connection")
        except Exception as e:
            print(f"✗ Error initializing database: {str(e)}")
            import traceback
            traceback.print_exc()
            print("\n⚠ Server will start but database operations may fail!")
            print("Please check your DATABASE_URL environment variable")
    
    print("\n" + "="*50)
    print("Starting Flask server on http://0.0.0.0:5002")
    print("="*50 + "\n")
    app.run(debug=True, host='0.0.0.0', port=5002)

